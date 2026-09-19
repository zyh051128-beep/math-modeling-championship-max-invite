"""Public synthetic DOCX/PDF deployment test, not a mathematical-modeling paper."""
import importlib.util
import json
from pathlib import Path
import shutil
import sys
from docx import Document

def run(skill_root, root):
    root.mkdir()
    code = root / 'code'
    code.mkdir()
    for filename in ('export_word_pdf.py', 'word_equations.py'):
        shutil.copyfile(skill_root / 'scripts' / filename, code / filename)
    def load(name):
        spec = importlib.util.spec_from_file_location(name, code / (name + '.py'))
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        spec.loader.exec_module(module)
        return module
    equations = load('word_equations')
    exporter = load('export_word_pdf')
    document = Document()
    document.add_heading('MAXx portable document verification', 0)
    document.add_paragraph('Synthetic deployment test. This text verifies an editable Word source and actual LibreOffice conversion, not scientific accuracy.')
    document.add_paragraph('The equation below uses native editable subscript structure. No page screenshot is embedded in this document.')
    equations.append_equation(document.add_paragraph(), {'type':'sub', 'base':{'type':'text','value':'x'}, 'sub':{'type':'text','value':'i'}})
    document.add_paragraph('External accounts, licensed applications, desktop login and complete contest papers are outside this smoke test.')
    source = root / 'source.docx'
    document.save(source)
    result = exporter.export(root, 'source.docx', 'rendered.pdf', 'conversion.json', engine='libreoffice', timeout=180)
    if result['status'] != 'PASS' or result['pdf']['page_count'] < 1:
        raise RuntimeError('Actual DOCX export failed.')
    print(json.dumps({'status':'PASS','engine':result['engine']['name'],'page_count':result['pdf']['page_count'],
                      'boundary':'Synthetic editable DOCX actually converted by LibreOffice; visual fidelity of future papers remains subject to per-formula and per-page audit.'}))

if __name__ == '__main__':
    run(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())
